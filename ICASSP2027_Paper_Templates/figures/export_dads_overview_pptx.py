"""Export the current Matplotlib overview as native editable PowerPoint objects.

Requires python-pptx, Matplotlib, NumPy and the existing renderer's fonts.
Example: PYTHONPATH=/tmp/dads-pptx-runtime MPLCONFIGDIR=/tmp/dads-mpl python3 <this file>
"""
from pathlib import Path
import re
import runpy

import numpy as np
from matplotlib.path import Path as MplPath
from matplotlib.patches import FancyArrowPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent
state = runpy.run_path(str(OUT / 'render_dads_overview.py'))
fig, ax = state['fig'], state['ax']
fig.canvas.draw()
renderer = fig.canvas.get_renderer()
prs = Presentation()
prs.slide_width = Inches(14)
prs.slide_height = Inches(14 * fig.get_figheight() / fig.get_figwidth())
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
scale = 14 / fig.get_figwidth()
pixel_w, pixel_h = fig.bbox.width, fig.bbox.height


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(key, str(value))
    return node


def rgb(color):
    return RGBColor(*(round(255 * float(c)) for c in color[:3]))


def to_emu(vertices):
    out = np.array(vertices, dtype=float).copy()
    out[:, 0] *= prs.slide_width / pixel_w
    out[:, 1] = (pixel_h - out[:, 1]) * prs.slide_height / pixel_h
    return np.rint(out).astype(np.int64)


def native_path(path, face, edge, linewidth, name, dashed=False, hatch=None):
    """Use DrawingML custom geometry to preserve editable Bezier paths."""
    segments = list(path.iter_segments(curves=True, simplify=False))
    points = [np.array(v).reshape(-1, 2) for v, code in segments if code != MplPath.CLOSEPOLY]
    if not points:
        return
    all_points = to_emu(np.concatenate(points))
    origin = all_points.min(axis=0)
    extent = np.maximum(all_points.max(axis=0) - origin, 1)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(origin[0]), int(origin[1]),
                                   int(extent[0]), int(extent[1]))
    shape.name = name
    # Default theme effects otherwise add shadows to every converted primitive.
    style = shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        shape._element.remove(style)
    sppr = shape._element.spPr
    sppr.remove(sppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom'))
    geom = element('a:custGeom')
    for tag in ['a:avLst', 'a:gdLst', 'a:ahLst', 'a:cxnLst']:
        geom.append(element(tag))
    geom.append(element('a:rect', l=0, t=0, r='r', b='b'))
    paths = element('a:pathLst')
    mapping = {MplPath.MOVETO: 'a:moveTo', MplPath.LINETO: 'a:lnTo',
               MplPath.CURVE3: 'a:quadBezTo', MplPath.CURVE4: 'a:cubicBezTo'}
    # Arrow shafts and arrowheads share an artist but must have distinct fills.
    subpaths = []
    for vertices, code in segments:
        if code == MplPath.MOVETO:
            subpaths.append([])
        subpaths[-1].append((vertices, code))
    for subpath in subpaths:
        geo_path = element('a:path', w=int(extent[0]), h=int(extent[1]))
        closed = any(code == MplPath.CLOSEPOLY for _, code in subpath)
        if not closed or face is None or face[3] == 0:
            geo_path.set('fill', 'none')
        for vertices, code in subpath:
            if code == MplPath.CLOSEPOLY:
                geo_path.append(element('a:close'))
                continue
            op = element(mapping[code])
            for point in to_emu(np.array(vertices).reshape(-1, 2)) - origin:
                op.append(element('a:pt', x=int(point[0]), y=int(point[1])))
            geo_path.append(op)
        paths.append(geo_path)
    geom.append(paths)
    sppr.insert(1, geom)
    if face is None or face[3] == 0:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(face)
    if edge is None or edge[3] == 0 or linewidth == 0:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(edge)
        shape.line.width = Pt(linewidth * scale)
        if dashed:
            shape.line._get_or_add_ln().append(element('a:prstDash', val='dash'))
    if hatch:
        # Explicit clipped diagonal lines survive both PowerPoint and Impress.
        vertices = np.concatenate(points)
        xmin, ymin = vertices.min(axis=0)
        xmax, ymax = vertices.max(axis=0)
        h = ymax - ymin
        hatch_vertices, codes = [], []
        for start in np.arange(xmin - h, xmax, 12):
            x1, x2 = max(xmin, start), min(xmax, start + h)
            if x1 < x2:
                hatch_vertices.extend([(x1, ymin + x1 - start), (x2, ymin + x2 - start)])
                codes.extend([MplPath.MOVETO, MplPath.LINETO])
        if hatch_vertices:
            native_path(MplPath(hatch_vertices, codes), None, edge, .45,
                        'Editable reference hatching')
    return shape


def formula_runs(value):
    """Keep formulas editable using Unicode symbols and native subscript runs."""
    for match in re.finditer(r'\$([^$]+)\$|([^$]+)', value):
        math, prose = match.groups()
        if prose is not None:
            yield prose, False
            continue
        math = math.replace(r'\widetilde E', 'E\u0303').replace(r'\hat y', 'y\u0302')
        for old, new in [(r'\Delta', 'Δ'), (r'\partial', '∂'), (r'\ell', 'ℓ'), (r'\odot', '⊙')]:
            math = math.replace(old, new)
        previous = 0
        for sub in re.finditer(r'_(?:\{([^}]+)\}|(.))', math):
            if sub.start() > previous:
                yield math[previous:sub.start()], False
            yield sub.group(1) or sub.group(2), True
            previous = sub.end()
        if previous < len(math):
            yield math[previous:], False


def native_text(artist):
    from matplotlib.colors import to_rgba
    text = artist.get_text()
    bbox = artist.get_window_extent(renderer)
    xy = artist.get_transform().transform(artist.get_position())
    center = to_emu([xy])[0]
    size = artist.get_fontsize() * scale
    width = max(bbox.width / pixel_w * prs.slide_width + Pt(size * .8), Pt(size * 1.2))
    height = Pt(size * 1.8)
    align = artist.get_ha()
    x = center[0] - (width / 2 if align == 'center' else width if align == 'right' else 0)
    shape = slide.shapes.add_textbox(int(x), int(center[1] - height / 2), int(width), int(height))
    shape.name = 'Text: ' + text.replace('$', '')
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = tf.paragraphs[0]
    paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    for content, subscript in formula_runs(text):
        for token in re.split('([⊙∂ℓ])', content):
            if not token:
                continue
            run = paragraph.add_run()
            run.text = token
            run.font.name = 'DejaVu Sans' if token in {'⊙', '∂', 'ℓ'} else 'Comic Sans MS'
            run.font.size = Pt(size)
            run.font.bold = artist.get_weight() == 'bold'
            run.font.color.rgb = rgb(to_rgba(artist.get_color()))
            if subscript:
                run._r.get_or_add_rPr().set('baseline', '-22000')


# Preserve visual stacking across fills, connectors, details, and labels.
artists = [(a.get_zorder(), i, a) for i, a in enumerate(ax.get_children())
           if a in ax.patches or a in ax.lines or a in ax.texts]
for _, _, artist in sorted(artists, key=lambda item: item[:2]):
    if artist in ax.texts:
        native_text(artist)
    elif artist in ax.patches:
        is_arrow = isinstance(artist, FancyArrowPatch)
        native_path(artist.get_path().transformed(artist.get_transform()),
                    artist.get_facecolor(), artist.get_edgecolor(), artist.get_linewidth(),
                    'Gray arrow' if is_arrow else 'Editable diagram shape',
                    dashed=is_arrow and artist.get_linestyle() != 'solid', hatch=artist.get_hatch())
    else:
        from matplotlib.colors import to_rgba
        native_path(artist.get_path().transformed(artist.get_transform()), None,
                    to_rgba(artist.get_color()), artist.get_linewidth(), 'Editable line',
                    dashed=artist.is_dashed())

prs.core_properties.title = 'DADS overview — editable Comic Sans version'
prs.core_properties.subject = 'Decision-aligned selection and frozen-head counterfactual recognition'
slide.notes_slide.notes_text_frame.text = (
    'All diagram objects and text are editable. Comic Sans MS is used for lettering. '
    'Formula subscripts are editable text runs; curves and arrows are editable freeform shapes. '
    'Feature glyphs are schematic. Source: render_dads_overview.py and export_dads_overview_pptx.py.'
)
target = OUT / 'dads_overview.pptx'
prs.save(target)
print(f'Saved {target}: {len(slide.shapes)} editable objects, {len(ax.texts)} text boxes.')
